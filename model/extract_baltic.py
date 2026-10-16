import openpyxl, csv
from datetime import datetime
SRC='/home/claude/fro/baltic_vlcc_suezmax_2022_2026.xlsx'
OUT='/home/claude/fm/models/frontline/data/baltic_weekly.csv'
wb=openpyxl.load_workbook(SRC, data_only=True)
v, s = wb['VLCC TCE'], wb['Suezmax TCE']

def grab(ws, cols):
    d={}
    for r in range(5, ws.max_row+1):
        dt=ws.cell(r,3).value
        if not isinstance(dt, datetime): continue
        d[dt.strftime('%Y-%m-%d')] = {name: ws.cell(r,c).value for name,c in cols.items()}
    return d
V = grab(v, {'TD3C':6,'TD15':7,'TD22':8,'TD34':9})
S = grab(s, {'TD20':6,'TD27':7,'TD33':8,'TD6':9})

weeks = sorted(set(V)|set(S))
FIELDS=['week_ending','TD3C','TD15','TD22','TD34','TD20','TD27','TD33','TD6','source']
rows=[]
for w in weeks:
    rec={'week_ending':w,'source':'Baltic Exchange weekly tanker report'}
    rec.update(V.get(w,{})); rec.update(S.get(w,{}))
    rows.append(rec)

# splice in the tracker weeks (they carry wk40/41 which post-date the workbook)
tr=openpyxl.load_workbook('/home/claude/fro/u9/t.xlsx')['Tracker']
R={6:'TD3C',7:'TD15',8:'TD22',9:'TD20',10:'TD27',11:'TD33'}
idx={r['week_ending']:r for r in rows}
added=0
for c in range(5, tr.max_column+1):
    d=tr.cell(4,c).value
    if not isinstance(d,datetime): continue
    k=d.strftime('%Y-%m-%d')
    tgt=idx.get(k)
    if tgt is None:
        tgt={'week_ending':k,'source':'Baltic Exchange weekly tanker report'}; rows.append(tgt); idx[k]=tgt; added+=1
    for r,name in R.items():
        val=tr.cell(r,c).value
        if val is not None: tgt[name]=val

rows.sort(key=lambda x:x['week_ending'])
with open(OUT,'w',newline='') as f:
    w=csv.DictWriter(f,fieldnames=FIELDS,extrasaction='ignore'); w.writeheader(); w.writerows(rows)

filled=sum(1 for r in rows if r.get('TD3C') is not None)
print(f'{len(rows)} weeks {rows[0]["week_ending"]} -> {rows[-1]["week_ending"]}   ({added} spliced from tracker)')
print(f'TD3C populated in {filled}; gaps: {len(rows)-filled}')
for r in rows[-3:]: print(' ', r['week_ending'], 'TD3C', r.get('TD3C'), 'TD20', r.get('TD20'))
