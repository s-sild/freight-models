import openpyxl, csv, os
from datetime import datetime
OUT='/home/claude/fm/models/frontline/data'
os.makedirs(OUT, exist_ok=True)

# ---------- Baltic weekly routes, from the tracker ----------
tw = openpyxl.load_workbook('/home/claude/fro/u9/t.xlsx')
tr = tw['Tracker']
ROUTES = {6:'TD3C',7:'TD15',8:'TD22',9:'TD20',10:'TD27',11:'TD33'}
rows=[]
for c in range(5, tr.max_column+1):
    d = tr.cell(4,c).value
    if not isinstance(d, datetime): continue
    rec = {'week_ending': d.strftime('%Y-%m-%d')}
    for r,name in ROUTES.items():
        rec[name] = tr.cell(r,c).value
    rec['source'] = 'Baltic Exchange weekly tanker report'
    rows.append(rec)
with open(f'{OUT}/baltic_weekly.csv','w',newline='') as f:
    w=csv.DictWriter(f, fieldnames=['week_ending']+list(ROUTES.values())+['source']); w.writeheader(); w.writerows(rows)
print(f'baltic_weekly.csv  {len(rows)} weeks, {rows[0]["week_ending"]} to {rows[-1]["week_ending"]}')

# ---------- FRO close ----------
px=[]
for c in range(5, tr.max_column+1):
    d=tr.cell(4,c).value; v=tr.cell(65,c).value
    if isinstance(d,datetime) and v is not None:
        px.append({'date':d.strftime('%Y-%m-%d'),'close':v,'source':'market close'})
with open(f'{OUT}/price.csv','w',newline='') as f:
    w=csv.DictWriter(f,fieldnames=['date','close','source']); w.writeheader(); w.writerows(px)
print(f'price.csv          {len(px)} observations')

# ---------- fixtures ----------
fx=tw['Fixtures']; out=[]
for r in range(2, fx.max_row+1):
    if not fx.cell(r,2).value: continue
    d=fx.cell(r,1).value
    out.append({'reported': d.strftime('%Y-%m-%d') if isinstance(d,datetime) else d,
        'vessel':fx.cell(r,2).value,'build_year':fx.cell(r,3).value,'operator':fx.cell(r,4).value,
        'charterer':fx.cell(r,5).value,'status':fx.cell(r,6).value,'open_port':fx.cell(r,7).value,
        'route':fx.cell(r,8).value,'load_mts':fx.cell(r,9).value,'rv_tce':fx.cell(r,10).value,
        'rv_days':fx.cell(r,11).value,'source':''})
with open(f'{OUT}/fro_fixtures.csv','w',newline='') as f:
    w=csv.DictWriter(f,fieldnames=list(out[0].keys())); w.writeheader(); w.writerows(out)
print(f'fro_fixtures.csv   {len(out)} fixtures')

# ---------- FRO reported quarterly, from the calibration workbook ----------
b=openpyxl.load_workbook('/home/claude/fro/u6/b.xlsx')['FRO reported']
q={}
for r in range(5,64):
    qt=b.cell(r,1).value; seg=b.cell(r,2).value
    if not qt or not seg or 'H' in str(qt): continue
    seg={'VLCC':'vlcc','Suezmax':'smax','LR2/Aframax':'lr2'}.get(seg)
    if not seg: continue
    d=q.setdefault(qt, {'quarter':qt,'source':b.cell(r,9).value})
    d[f'{seg}_spot_tce']=b.cell(r,3).value
    d[f'{seg}_breakeven']=b.cell(r,4).value
    d[f'{seg}_spot_days']=b.cell(r,5).value
    d[f'{seg}_next_contracted']=b.cell(r,6).value
    d[f'{seg}_next_cov_pct']=b.cell(r,7).value
# income-statement items we pulled from the reports
extra={'2025Q3':(82830,42505,13571),'2025Q4':(82199,230390,11803),
       '2026Q1':(75996,344922,26993),'2026Q2':(71338,580210,83384)}
for k,(dep,adj,tc) in extra.items():
    if k in q: q[k].update({'depreciation_kusd':dep,'adjusted_profit_kusd':adj,'tc_tce_kusd':tc})
cols=['quarter','vlcc_spot_tce','smax_spot_tce','lr2_spot_tce','vlcc_spot_days','smax_spot_days','lr2_spot_days',
      'vlcc_breakeven','smax_breakeven','lr2_breakeven','vlcc_next_contracted','vlcc_next_cov_pct',
      'smax_next_contracted','smax_next_cov_pct','lr2_next_contracted','lr2_next_cov_pct',
      'tc_tce_kusd','depreciation_kusd','adjusted_profit_kusd','source']
with open(f'{OUT}/fro_quarterly.csv','w',newline='') as f:
    w=csv.DictWriter(f,fieldnames=cols,extrasaction='ignore'); w.writeheader()
    for k in sorted(q): w.writerow(q[k])
print(f'fro_quarterly.csv  {len(q)} quarters, {min(q)} to {max(q)}')

# ---------- Fearnleys 1-yr T/C (only this week is known) ----------
with open(f'{OUT}/tc_weekly.csv','w',newline='') as f:
    w=csv.writer(f); w.writerow(['week_ending','vlcc_1yr_tc','smax_1yr_tc','afra_1yr_tc','vlcc_fixtures','vlcc_meg_avail_30d','source'])
    w.writerow(['2026-10-15',225000,165000,100000,50,99,'Fearnleys weekly'])
print('tc_weekly.csv      1 week (series starts here)')
